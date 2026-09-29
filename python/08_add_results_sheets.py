"""08 - Results workbook, part 2: adds the Accuracy, Matching, Matched pairs, DiD results,
Event study, Fragmentation and Hotspots sheets to results/Vegetation_loss_results.xlsx.
Needs the outputs of 03, 04, 05, 06 and 07.
"""
import sys, os; sys.path.insert(0, os.path.dirname(__file__))
from paths import RES
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import ScatterChart, Reference, Series
F='Arial'
HDR=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E79'); B=Font(name=F,bold=True); N=Font(name=F); I=Font(name=F,italic=True)
GOOD=PatternFill('solid',fgColor='E2EFDA'); BAD=PatternFill('solid',fgColor='FCE4D6')
wb=openpyxl.load_workbook(RES/'Vegetation_loss_results.xlsx')
for n in ['Accuracy','Matching','Matched pairs','DiD results','Event study','Fragmentation','Hotspots','Paper tables plan']:
    if n in wb.sheetnames: del wb[n]
def header(ws,r,vals):
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=HDR; c.fill=HF; c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
# Matching
bal=pd.read_csv(RES/'matching_balance.csv'); ws=wb.create_sheet('Matching')
ws['A1']='Step 9 — Matching of control segments (covariate balance)'; ws['A1'].font=Font(name=F,bold=True,size=13)
ws['A2']='Method: 2 nearest control segments per Section 1 segment, Mahalanobis distance on 2020 vegetation and built-up shares, with replacement, caliper 0.5. Section 1 segments with no control within the caliper are dropped (common support).'
ws['A3']='Result: 22 of 47 Section 1 segments matched to 21 distinct control segments. SMD = standardised mean difference; |SMD| < 0.25 = acceptable balance (Stuart, 2010).'
for c in ('A2','A3'): ws[c].font=I
header(ws,5,['Covariate (2020 baseline)','Used in matching','Section 1 mean (all 47)','Control mean (all 110)','SMD before','Section 1 mean (matched)','Control mean (matched, weighted)','SMD after'])
for i,r in enumerate(bal.itertuples(index=False),6):
    vals=list(r)
    for j,v in enumerate(vals,1):
        c=ws.cell(i,j,v); c.font=N
        if j in (3,4,6,7): c.number_format='0.00'
        if j in (5,8):
            c.number_format='0.00'; c.fill=GOOD if abs(v)<0.25 else BAD
i+=2
ws.cell(i,1,'Interpretation: matching removes the large urbanisation imbalance (vegetation SMD −1.36 → −0.03; built-up +1.91 → +0.06). Matched Section 1 segments are the less urbanised, eastern part of the corridor. Elevation and water share remain imbalanced; specification C in "DiD results" adjusts for these baseline differences.').font=I
ws.cell(i,1).alignment=Alignment(wrap_text=True); ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=8); ws.row_dimensions[i].height=45
for c,wd in zip('ABCDEFGH',[34,11,14,14,10,14,16,10]): ws.column_dimensions[c].width=wd
ws.row_dimensions[5].height=45
# Matched pairs
p=pd.read_csv(RES/'matched_pairs.csv'); wp=wb.create_sheet('Matched pairs')
header(wp,1,['Section 1 segment','Matched control segment','Mahalanobis distance'])
for i,r in enumerate(p.itertuples(index=False),2):
    for j,v in enumerate(r,1): wp.cell(i,j,v).font=N
    wp.cell(i,3).number_format='0.000'
for c,wd in zip('ABC',[18,28,20]): wp.column_dimensions[c].width=wd
# DiD results
res=pd.read_csv(RES/'did_results.csv'); wr=wb.create_sheet('DiD results')
wr['A1']='Pooled difference-in-differences: extra % of 2020 vegetation lost in Section 1 after construction (DS2025–26 vs DS2021–24)'; wr['A1'].font=Font(name=F,bold=True,size=13)
wr['A2']='Unit and year fixed effects; standard errors clustered by 1 km segment; wild cluster bootstrap (Rademacher, 499 draws, null imposed) for the matched specification. Values computed in Python (statsmodels).'; wr['A2'].font=I
header(wr,4,['Zone','Distance','Specification','Estimate (pp)','95% CI low','95% CI high','p (clustered)','p (wild bootstrap)','Section 1 units','Control units'])
for i,r in enumerate(res.itertuples(index=False),5):
    for j,v in enumerate(r,1):
        c=wr.cell(i,j,None if (isinstance(v,float) and pd.isna(v)) else v); c.font=B if 'Matched' in r.spec else N
        if j in (4,5,6): c.number_format='0.0'
        if j in (7,8): c.number_format='0.000'
for c,wd in zip('ABCDEFGHIJ',[7,12,40,12,10,10,11,12,10,10]): wr.column_dimensions[c].width=wd
wr.row_dimensions[4].height=32
# Event study
es=pd.read_csv(RES/'event_study.csv'); we=wb.create_sheet('Event study')
we['A1']='Event study (matched sample): difference Section 1 − controls in % vegetation lost, relative to DS2024'; we['A1'].font=Font(name=F,bold=True,size=13)
we['A2']='Pre-construction values (2021–2023) should be near 0 if trends were parallel. Values in percentage points; 95% CI clustered by segment.'; we['A2'].font=I
header(we,4,['Year']+[f'{z} est' for z in ['F','B1','B2','B3']]+[f'{z} low' for z in ['F','B1','B2','B3']]+[f'{z} high' for z in ['F','B1','B2','B3']])
m=es[es.spec.str.startswith('B')]
for i,y in enumerate(range(2021,2027),5):
    we.cell(i,1,y).font=N
    for k,z in enumerate(['F','B1','B2','B3']):
        r=m[(m.zone==z)&(m.year==y)].iloc[0]
        for off,col in [(0,'estimate'),(4,'ci_low'),(8,'ci_high')]:
            c=we.cell(i,2+k+off,round(r[col],2)); c.number_format='0.0'; c.font=N
for c in 'ABCDEFGHIJKLM': we.column_dimensions[c].width=9
ch=ScatterChart(); ch.title='Event study, matched sample (pp vs DS2024)'; ch.style=13; ch.height=9; ch.width=18
ch.x_axis.title='Dry season'; ch.y_axis.title='Section 1 − control (pp)'; ch.x_axis.scaling.min=2020; ch.x_axis.scaling.max=2027; ch.x_axis.majorUnit=1; ch.x_axis.number_format='0'; ch.x_axis.crosses='min'; ch.legend.position='b'
xs=Reference(we,min_col=1,min_row=5,max_row=10)
cols={'F':'7F7F7F','B1':'C00000','B2':'ED7D31','B3':'2F5597'}
for k,z in enumerate(['F','B1','B2','B3']):
    sr=Series(Reference(we,min_col=2+k,min_row=4,max_row=10),xs,title_from_data=True)
    sr.marker.symbol='circle'; sr.smooth=False; sr.graphicalProperties.line.solidFill=cols[z]; sr.marker.graphicalProperties.solidFill=cols[z]
    ch.series.append(sr)
ch.x_axis.delete=False; ch.y_axis.delete=False
we.add_chart(ch,'A13')
# Accuracy
acc=pd.read_csv(RES/'accuracy_summary.csv'); mat=pd.read_csv(RES/'accuracy_error_matrices.csv')
wa=wb.create_sheet('Accuracy',1)
wa['A1']='Accuracy assessment (300 stratified random points; Olofsson et al., 2014)'; wa['A1'].font=Font(name=F,bold=True,size=13)
wa['A2']='Accuracies and areas are area-weighted estimates; ± = 95% confidence interval. Reference labels from Esri Wayback (2020) and Wayback/PlanetScope (2025–26).'; wa['A2'].font=I
r=4
for name,g in acc.groupby('map',sort=False):
    wa.cell(r,1,f'{name} — overall accuracy {100*g.overall.iloc[0]:.1f}% ± {100*g.overall_ci.iloc[0]:.1f}').font=B; r+=1
    header(wa,r,['Class',"User's (%)",'± CI',"Producer's (%)",'± CI','Mapped area (ha)','Error-adjusted area (ha)','± CI (ha)']); r+=1
    for x in g.itertuples():
        for j,v in enumerate([x[9],100*x.users,100*x.users_ci,100*x.producers,100*x.producers_ci,x.map_area_ha,x.adj_area_ha,x.adj_area_ci],1):
            c=wa.cell(r,j,v); c.font=N; c.number_format='0.0' if 1<j<6 else '#,##0'
        r+=1
    m=mat[mat['map']==name].drop(columns='map').dropna(axis=1,how='all')
    wa.cell(r,1,'Sample counts (rows = map, columns = reference)').font=I; r+=1
    header(wa,r,list(m.columns)); r+=1
    for row in m.itertuples(index=False):
        for j,v in enumerate(row,1): wa.cell(r,j,v).font=N
        r+=1
    r+=1
for c,wd in zip('ABCDEFGH',[24,11,8,13,8,15,19,11]): wa.column_dimensions[c].width=wd
# Fragmentation
fs=pd.read_csv(RES/'fragmentation_summary.csv'); fd=pd.read_csv(RES/'fragmentation_did.csv')
wf=wb.create_sheet('Fragmentation')
wf['A1']='Vegetation fragmentation, matched sample (8-neighbour patches; patches < 0.1 ha removed)'; wf['A1'].font=Font(name=F,bold=True,size=13)
wf['A2']='Landscape = one segment\'s corridor (footprint + 0–500 m), 500 m–1 km or 1–5 km band. PLAND % vegetation; PD patches/100 ha; MPA mean patch area (ha); ED edge density (m/ha); LPI largest patch (% of land). Means weighted by matching weights.'; wf['A2'].font=I
header(wf,4,list(fs.columns))
for i,row in enumerate(fs.itertuples(index=False),5):
    for j,v in enumerate(row,1):
        c=wf.cell(i,j,v); c.font=N
        if j>4: c.number_format='0.00'
r=i+2
wf.cell(r,1,'Difference-in-differences on each metric (unit + year fixed effects, matched weights, SE clustered by segment; reference year DS2024)').font=B; r+=1
header(wf,r,list(fd.columns)); r+=1
for row in fd.itertuples(index=False):
    for j,v in enumerate(row,1):
        c=wf.cell(r,j,v); c.font=B if ('Construction' in str(row.term) and row.p<0.05) else N
        if j in (4,5,6): c.number_format='0.00'
        if j==7: c.number_format='0.000'
    r+=1
for c,wd in zip('ABCDEFGHI',[12,34,14,12,11,11,11,11,11]): wf.column_dimensions[c].width=wd
# Hotspots
hs=pd.read_csv(RES/'hotspot_summary.csv'); wh=wb.create_sheet('Hotspots')
wh['A1']='Hotspots of persistent vegetation loss in Section 1 (Getis-Ord Gi*, 999 permutations, FDR q < 0.05)'; wh['A1'].font=Font(name=F,bold=True,size=13)
wh['A2']='Persistent loss: vegetated at the start and non-vegetated in the last two years of the period. Sensitivity to hexagon size (ha) and distance band (m). The paper map uses 5 ha / 500 m.'; wh['A2'].font=I
header(wh,4,['Hexagon (ha)','Distance band (m)','Period','Hexagons analysed','Hot spots','Cold spots','Hot spots within 500 m of road (%)','All hexagons within 500 m (%)','Median distance of hot spots to road (m)','Persistent loss (ha)'])
for i,row in enumerate(hs.itertuples(index=False),5):
    for j,v in enumerate(row,1):
        c=wh.cell(i,j,v); c.font=B if (row.hex_ha==5 and row.band_m==500) else N
        if j>6: c.number_format='0.0'
for c,wd in zip('ABCDEFGHIJ',[11,11,32,12,10,10,16,16,18,14]): wh.column_dimensions[c].width=wd
wh.row_dimensions[4].height=45
wb.save(RES/'Vegetation_loss_results.xlsx'); print(wb.sheetnames)
